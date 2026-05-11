"""Generate a simple proposal DOCX that follows the local opening-report style."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


REPO_ROOT = Path(__file__).resolve().parents[2]
HARNESS_DIR = REPO_ROOT / "experiments" / "harness"
SOURCE_PATH = HARNESS_DIR / "code2workspace_开题报告.md"
OUTPUT_PATH = HARNESS_DIR / "code2workspace_开题报告.docx"


def set_run_font(run, *, east_asia: str = "宋体", ascii_name: str = "Times New Roman", size: int = 12, bold: bool | None = None) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clean_text(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()


def add_body_paragraph(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(clean_text(text))
    set_run_font(run, size=12)


def add_heading(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(clean_text(text))
    set_run_font(run, east_asia="黑体", size=14, bold=True)


def add_numbered_item(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(24)
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run(clean_text(text))
    set_run_font(run, size=12)


def parse_sections(lines: list[str]) -> tuple[str, dict[str, list[str]]]:
    title = ""
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for raw in lines:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("## 题目"):
            current = "题目"
            sections[current] = []
            continue
        if line.startswith("## "):
            current = line[3:].strip()
            sections[current] = []
            continue
        if current is None:
            continue
        sections[current].append(line)
    title_lines = sections.get("题目", [])
    if title_lines:
        title = clean_text(title_lines[0])
    return title, sections


def add_cover(doc: Document, title: str) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("本科毕业设计（论文）开题报告")
    set_run_font(run, east_asia="黑体", size=18, bold=True)

    for _ in range(2):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("论文题目")
    set_run_font(run, east_asia="宋体", size=14, bold=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, east_asia="黑体", size=18, bold=True)

    doc.add_page_break()


def parse_schedule_rows(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        if not line.startswith("|"):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if len(cells) != 3:
            continue
        if cells[0] == "序号" or set("".join(cells)) <= {"-", " "}:
            continue
        rows.append(cells)
    return rows


def build_doc() -> Document:
    lines = SOURCE_PATH.read_text(encoding="utf-8").splitlines()
    title, sections = parse_sections(lines)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Pt(85)
    section.bottom_margin = Pt(85)
    section.left_margin = Pt(90)
    section.right_margin = Pt(90)

    add_cover(doc, title)

    ordered_sections = [
        "一、课题背景及意义（含国内外研究现状综述）",
        "二、课题研究主要内容",
        "三、研究方案和思路",
        "四、论文框架结构",
        "五、参考文献",
        "六、工作进度安排",
    ]

    for name in ordered_sections:
        add_heading(doc, name)
        lines = sections.get(name, [])
        if name == "六、工作进度安排":
            rows = parse_schedule_rows(lines)
            table = doc.add_table(rows=1, cols=3)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            headers = ["序号", "设计（论文）各阶段任务", "时间安排"]
            for idx, value in enumerate(headers):
                cell = table.rows[0].cells[idx]
                cell.text = value
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        set_run_font(run, size=11, bold=True)
            for row in rows:
                cells = table.add_row().cells
                for idx, value in enumerate(row):
                    cells[idx].text = clean_text(value)
                    for para in cells[idx].paragraphs:
                        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in para.runs:
                            set_run_font(run, size=11)
            continue

        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith("### "):
                add_body_paragraph(doc, stripped[4:])
            elif re.match(r"^\d+\.", stripped):
                add_numbered_item(doc, stripped)
            elif stripped.startswith("[") and "]" in stripped:
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.5
                run = paragraph.add_run(clean_text(stripped))
                set_run_font(run, size=11)
            else:
                add_body_paragraph(doc, stripped)
    return doc


def main() -> None:
    doc = build_doc()
    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
