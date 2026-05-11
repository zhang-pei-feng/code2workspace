"""Generate a simple task-book DOCX following the local sample style."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


REPO_ROOT = Path(__file__).resolve().parents[2]
HARNESS_DIR = REPO_ROOT / "experiments" / "harness"
SOURCE_PATH = HARNESS_DIR / "code2workspace_任务书.md"
OUTPUT_PATH = HARNESS_DIR / "code2workspace_任务书.docx"


def set_run_font(run, *, east_asia: str = "宋体", ascii_name: str = "Times New Roman", size: int = 12, bold: bool | None = None) -> None:
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:ascii"), ascii_name)
    rfonts.set(qn("w:hAnsi"), ascii_name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def clean_text(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    set_run_font(run, east_asia="黑体", size=18, bold=True)


def add_body(doc: Document, text: str, *, indent: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(clean_text(text))
    set_run_font(run, size=12)


def build_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(85)
    section.bottom_margin = Pt(85)
    section.left_margin = Pt(90)
    section.right_margin = Pt(90)

    lines = SOURCE_PATH.read_text(encoding="utf-8").splitlines()
    add_title(doc, "本科毕业设计（论文）任务书")

    for raw in lines[1:]:
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("#"):
            continue
        if re.match(r"^\d+、", line):
            add_body(doc, line)
        elif line.startswith("（") and "）" in line[:4]:
            add_body(doc, line)
        else:
            add_body(doc, line, indent=False)
    return doc


def main() -> None:
    doc = build_doc()
    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
