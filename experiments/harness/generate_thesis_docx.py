"""Generate a thesis DOCX that follows the local undergraduate sample layout."""

from __future__ import annotations

from pathlib import Path
import re
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parents[2]
HARNESS_DIR = REPO_ROOT / "experiments" / "harness"
TEMPLATE_PATH = Path("/home/zhangpf/参考模版（注意内容无关）/汪峻峰-毕业论文.docx")
SOURCE_PATH = HARNESS_DIR / "THESIS_FULL_DRAFT_ZH.md"
OUTPUT_PATH = HARNESS_DIR / "code2workspace_毕业论文.docx"

CHAPTER_HEADER_MAP = {
    "第一章 绪论": "第一章 绪论",
    "第二章 相关技术与理论基础": "第二章 相关技术与理论基础",
    "第三章 code2workspace 系统设计与实现": "第三章 code2workspace 系统设计与实现",
    "第四章 基于 Surface 的 Harness 优化方法": "第四章 基于 Surface 的 Harness 优化方法",
    "第五章 实验结果与分析": "第五章 实验结果与分析",
}

CHAPTER_BODY_TITLE_MAP = {
    "第一章 绪论": "绪论",
    "第二章 相关技术与理论基础": "相关技术与理论基础",
    "第三章 code2workspace 系统设计与实现": "code2workspace 系统设计与实现",
    "第四章 基于 Surface 的 Harness 优化方法": "基于 Surface 的 Harness 优化方法",
    "第五章 实验结果与分析": "实验结果与分析",
}


def remove_all_body_content(doc: Document) -> None:
    """Clear document body while keeping styles available from the template."""
    body = doc._element.body
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            continue
        body.remove(child)


def clean_inline_markdown(text: str) -> str:
    """Remove lightweight markdown markers for docx output."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("`", "")
    text = text.replace("**", "")
    text = text.replace("__", "")
    return text.strip()


def set_run_font(run, *, east_asia: str | None = None, ascii_name: str | None = None, size: int | None = None, bold: bool | None = None) -> None:
    """Apply explicit run font settings when style defaults are not enough."""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    if east_asia is not None:
        rfonts.set(qn("w:eastAsia"), east_asia)
    if ascii_name is not None:
        rfonts.set(qn("w:ascii"), ascii_name)
        rfonts.set(qn("w:hAnsi"), ascii_name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_field(paragraph, instruction: str) -> None:
    """Insert a Word field."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = ""
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def set_page_number_start(section, start: int) -> None:
    """Set page number restart for one section."""
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def clear_hdrftr(container) -> None:
    """Remove all children from a header/footer container."""
    element = container._element
    for child in list(element):
        element.remove(child)


def set_section_header_footer(section, header_text: str, *, page_number: bool = True) -> None:
    """Set one section header/footer."""
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header
    clear_hdrftr(header)
    p = header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(header_text.replace("`", ""))
    set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size=10)

    footer = section.footer
    clear_hdrftr(footer)
    fp = footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if page_number:
        add_field(fp, "PAGE")


def parse_markdown_tables(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    """Parse a simple markdown table from lines[start:]."""
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        raw = lines[index].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        rows.append(cells)
        index += 1
    return rows, index


def apply_normal_paragraph_format(paragraph) -> None:
    """Apply the sample thesis body formatting."""
    paragraph.style = "Normal"
    paragraph.paragraph_format.first_line_indent = Pt(24)
    paragraph.paragraph_format.line_spacing = 1.5


def add_body_paragraph(doc: Document, text: str) -> None:
    """Add one standard body paragraph."""
    p = doc.add_paragraph(style="Normal")
    apply_normal_paragraph_format(p)
    run = p.add_run(clean_inline_markdown(text))
    set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size=12)


def add_list_paragraph(doc: Document, text: str) -> None:
    """Add one simple list paragraph."""
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(clean_inline_markdown(text))
    set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size=12)


def add_heading(doc: Document, level: int, text: str) -> None:
    """Add one heading paragraph using the template heading styles."""
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_paragraph(style=style)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(clean_inline_markdown(text))
    if level == 1:
        set_run_font(run, east_asia="黑体", ascii_name="Times New Roman", size=16, bold=True)
    elif level == 2:
        set_run_font(run, east_asia="黑体", ascii_name="Times New Roman", size=14, bold=True)
    else:
        set_run_font(run, east_asia="黑体", ascii_name="Times New Roman", size=13, bold=True)


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    """Render a markdown table as a Word table."""
    if len(rows) < 2:
        return
    header = rows[0]
    data_rows = rows[2:] if len(rows) >= 2 and all(set(cell) <= {"-", ":"} for cell in rows[1]) else rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Normal Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for idx, value in enumerate(header):
        hdr_cells[idx].text = clean_inline_markdown(value)
        for para in hdr_cells[idx].paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row in data_rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = clean_inline_markdown(value)
            for para in cells[idx].paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def extract_between(text: str, start: str, end: str) -> str:
    """Return the text between two markers."""
    left = text.index(start) + len(start)
    right = text.index(end, left)
    return text[left:right].strip()


def add_centered_paragraph(doc: Document, text: str, *, size: int = 12, bold: bool = False, east_asia: str = "宋体") -> None:
    """Add one centered paragraph."""
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(clean_inline_markdown(text))
    set_run_font(run, east_asia=east_asia, ascii_name="Times New Roman", size=size, bold=bold)


def build_cover(doc: Document) -> None:
    """Build the cover page based on the sample layout."""
    doc.add_paragraph()
    add_centered_paragraph(doc, "本科毕业设计（论文）", size=18, bold=True, east_asia="黑体")
    doc.add_paragraph()
    title = doc.add_paragraph(style="Normal")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("面向真实软件仓库任务的智能体工作空间系统设计与实现")
    set_run_font(run, east_asia="黑体", ascii_name="Times New Roman", size=20, bold=True)
    doc.add_paragraph()
    doc.add_paragraph()

    template_doc = Document(TEMPLATE_PATH)
    template_table = template_doc.tables[0]
    new_tbl = deepcopy(template_table._tbl)
    doc._body._element.append(new_tbl)
    table = doc.tables[-1]
    replacement_values = [
        "______________________________",
        "______________________________",
        "______________________________",
        "______________________________",
        "______________________________",
        "   年   月   日",
    ]
    for row_index, value in enumerate(replacement_values):
        table.cell(row_index, 1).text = value
    for cell in table._cells:
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size=15, bold=True)

    doc.add_paragraph()
    add_centered_paragraph(doc, "华南理工大学", size=18, bold=True, east_asia="宋体")


def build_statement_pages(doc: Document) -> None:
    """Add originality and authorization pages."""
    doc.add_page_break()
    add_centered_paragraph(doc, "学位论文原创性声明", size=16, bold=True, east_asia="宋体")
    add_body_paragraph(
        doc,
        "本人郑重声明：所呈交的论文是本人在导师的指导下独立进行研究所取得的成果。除了文中特别加以标注引用的内容外，本论文不包含任何其他个人或集体已经发表或撰写过的研究成果。对本文研究做出重要贡献的个人和集体，均已在文中以明确方式标明。本人完全意识到本声明的法律后果由本人承担。",
    )
    add_body_paragraph(doc, "作者签名：                日期：    年   月   日")
    doc.add_paragraph()
    add_centered_paragraph(doc, "学位论文版权使用授权书", size=16, bold=True, east_asia="宋体")
    add_body_paragraph(
        doc,
        "本人完全了解学校关于保留、使用学位论文的有关规定，同意学校保留并向有关部门或机构送交论文的复印件和电子版，允许论文被查阅和借阅；同意学校可以公布论文的全部或部分内容，并采用影印、缩印或其他复制手段保存和汇编本论文。",
    )
    add_body_paragraph(doc, "作者签名：                        日期：")
    add_body_paragraph(doc, "指导教师签名：                    日期：")
    add_body_paragraph(doc, "作者联系电话：                    电子邮箱：")


def build_abstract_pages(doc: Document, source_text: str) -> None:
    """Add Chinese and English abstracts."""
    zh = extract_between(source_text, "## 摘要", "## Abstract")
    en = extract_between(source_text, "## Abstract", "---\n\n## 目录")
    for title, block in [("摘  要", zh), ("Abstract", en)]:
        doc.add_page_break()
        add_heading(doc, 1, title)
        for para in [item.strip() for item in block.split("\n\n") if item.strip()]:
            add_body_paragraph(doc, para)


def build_toc_page(doc: Document) -> None:
    """Add a TOC page."""
    doc.add_page_break()
    add_heading(doc, 1, "目  录")
    p = doc.add_paragraph(style="Normal")
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u')


def build_from_markdown(doc: Document, source_text: str) -> None:
    """Render the main matter from markdown."""
    main_text = source_text.split("# 第一章 绪论", 1)[1]
    lines = ["# 第一章 绪论"] + main_text.splitlines()
    index = 0
    first_main_section = True
    current_section = None
    last_heading = ""
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        if stripped == "---":
            index += 1
            continue
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip().replace("`", "")
            if text.startswith("## "):
                text = text[3:].strip()
            if level == 1:
                text = CHAPTER_HEADER_MAP.get(text, text)
                if first_main_section:
                    current_section = doc.add_section(WD_SECTION.NEW_PAGE)
                    set_page_number_start(current_section, 1)
                    first_main_section = False
                else:
                    current_section = doc.add_section(WD_SECTION.NEW_PAGE)
                set_section_header_footer(current_section, text, page_number=True)
                body_heading_text = CHAPTER_BODY_TITLE_MAP.get(text, text)
                body_heading_text = re.sub(r"^第[一二三四五六七八九十]+章\s*", "", body_heading_text)
                add_heading(doc, 1, body_heading_text)
                last_heading = body_heading_text
                index += 1
                continue
            if level >= 1 and text.startswith("参考文献"):
                if not first_main_section:
                    current_section = doc.add_section(WD_SECTION.NEW_PAGE)
                    set_section_header_footer(current_section, "参考文献", page_number=True)
                add_heading(doc, 1, "参考文献")
                last_heading = "参考文献"
                index += 1
                continue
            if level >= 1 and text == "致谢":
                current_section = doc.add_section(WD_SECTION.NEW_PAGE)
                set_section_header_footer(current_section, "致谢", page_number=True)
                add_heading(doc, 1, "致谢")
                last_heading = "致谢"
                index += 1
                continue
            add_heading(doc, min(level, 3), text)
            last_heading = text
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, index = parse_markdown_tables(lines, index)
            add_markdown_table(doc, rows)
            continue
        if stripped.startswith("```"):
            language = stripped.strip("`").strip()
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip("\n"))
                index += 1
            if language == "mermaid":
                placeholder = doc.add_paragraph(style="Normal")
                placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = placeholder.add_run("（此处插入对应示意图，建议依据文中说明使用 draw.io 或 Visio 重绘）")
                set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size=12)
            else:
                for code_line in code_lines:
                    p = doc.add_paragraph(style="Normal")
                    p.paragraph_format.first_line_indent = Pt(0)
                    run = p.add_run(code_line)
                    set_run_font(run, east_asia="等线", ascii_name="Consolas", size=10)
            index += 1
            continue
        if stripped.startswith("- "):
            add_list_paragraph(doc, stripped)
            index += 1
            continue
        if re.match(r"^\d+\.\s", stripped):
            add_list_paragraph(doc, stripped)
            index += 1
            continue
        if stripped.startswith("> "):
            p = doc.add_paragraph(style="Normal")
            p.paragraph_format.first_line_indent = Pt(0)
            run = p.add_run(clean_inline_markdown(stripped[2:].strip()))
            set_run_font(run, east_asia="宋体", ascii_name="Times New Roman", size=12)
            index += 1
            continue
        add_body_paragraph(doc, stripped)
        index += 1


def main() -> None:
    """Generate the thesis docx file."""
    source_text = SOURCE_PATH.read_text(encoding="utf-8")
    doc = Document(TEMPLATE_PATH)
    remove_all_body_content(doc)
    for section in doc.sections:
        section.top_margin = doc.sections[0].top_margin
        section.bottom_margin = doc.sections[0].bottom_margin
        section.left_margin = doc.sections[0].left_margin
        section.right_margin = doc.sections[0].right_margin
    doc.sections[0].header.is_linked_to_previous = False
    doc.sections[0].footer.is_linked_to_previous = False
    clear_hdrftr(doc.sections[0].header)
    doc.sections[0].header.add_paragraph()
    clear_hdrftr(doc.sections[0].footer)
    doc.sections[0].footer.add_paragraph()

    build_cover(doc)
    build_statement_pages(doc)
    build_abstract_pages(doc, source_text)
    build_toc_page(doc)
    build_from_markdown(doc, source_text)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
